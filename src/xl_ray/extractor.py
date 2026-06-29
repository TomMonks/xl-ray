from pathlib import Path
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.formula import Tokenizer
from openpyxl.utils import range_boundaries, get_column_letter, coordinate_to_tuple

from .schema import WorkbookMetadata, VBAModule, NamedRange, ExcelTable, WorksheetData, CellData

from oletools.olevba import VBA_Parser
import re

def extract_precedents_from_formula(formula_str: str) -> list[str]:
    """
    Parses an Excel formula and returns a list of referenced cell addresses, 
    ranges, or named ranges.
    """
    if not formula_str or not formula_str.startswith('='):
        return []
        
    precedents = set()
    try:
        tok = Tokenizer(formula_str)
        for t in tok.items:
            # Tokenizer identifies ranges (A1, A1:B2, 'Sheet1'!A1) and names
            if t.type == 'OPERAND' and t.subtype == 'RANGE':
                # Strip out absolute reference markers ($) to keep things clean
                clean_ref = t.value.replace('$', '')
                precedents.add(clean_ref)
    except Exception:
        # Fail gracefully on extremely complex or malformed formulas
        pass
        
    return sorted(list(precedents))




def extract_worksheets(wb_formulas, wb_values) -> dict[str, WorksheetData]:
    worksheets_data = {}
    
    # --- OPTIMIZED: Track references without unrolling ranges ---
    exact_referenced_cells = set()  # Fast O(1) lookup for single cells like "Sheet1!A1"
    range_references = {}           # Dict mapping sheet_name -> list of (min_col, min_row, max_col, max_row)
    # ------------------------------------------------------------

    for sheet_f, sheet_v in zip(wb_formulas.worksheets, wb_values.worksheets):
        sheet_name = sheet_f.title
        state_map = {"visible": "Visible", "hidden": "Hidden", "veryHidden": "VeryHidden"}
        visibility = state_map.get(sheet_f.sheet_state, "Visible")
        
        cells_dict = {}
        sheet_tables = extract_tables(sheet_f)
        
        for row_f, row_v in zip(sheet_f.iter_rows(values_only=False), sheet_v.iter_rows(values_only=False)):
            for cell_f, cell_v in zip(row_f, row_v):
                if cell_f.value is None and cell_v.value is None:
                    continue
                    
                address = cell_f.coordinate
                formula_str = None
                is_array = False
                array_range = None
                precedents = []
                
                if cell_f.data_type == 'f':
                    val_f = cell_f.value
                    if isinstance(val_f, str):
                        formula_str = val_f
                    elif isinstance(val_f, ArrayFormula):
                        is_array = True
                        formula_str = str(getattr(val_f, "text", None) or "")
                        array_range = str(getattr(val_f, "ref", None) or "")
                    else:
                        formula_str = str(getattr(val_f, "value", None) or val_f)
                        
                    precedents = extract_precedents_from_formula(formula_str)
                    
                    # --- OPTIMIZED: Store ranges as mathematical boundaries ---
                    for p in precedents:
                        if "!" in p:
                            sheet_part, cell_part = p.split("!")
                            sheet_part = sheet_part.replace("'", "")
                        else:
                            sheet_part = sheet_name
                            cell_part = p
                            
                        if ":" in cell_part:
                            try:
                                bounds = range_boundaries(cell_part)
                                if sheet_part not in range_references:
                                    range_references[sheet_part] = []
                                range_references[sheet_part].append(bounds)
                            except Exception:
                                exact_referenced_cells.add(f"{sheet_part}!{cell_part}")
                        else:
                            exact_referenced_cells.add(f"{sheet_part}!{cell_part}")
                    # -----------------------------------------------------------
                    
                data_type = cell_v.data_type if cell_v.data_type else cell_f.data_type
                
                cells_dict[address] = CellData(
                    address=address,
                    value=cell_v.value,
                    formula=formula_str,
                    precedents=precedents,
                    is_array_formula=is_array,
                    array_range=array_range,
                    data_type=data_type
                )
                
        # --- Array Formula Spill Propagation ---
        array_parents = {addr: cell for addr, cell in cells_dict.items() if cell.is_array_formula and cell.array_range}

        for parent_addr, parent_cell in array_parents.items():
            if ':' in parent_cell.array_range:
                min_col, min_row, max_col, max_row = range_boundaries(parent_cell.array_range)

                for row in range(min_row, max_row + 1):
                    for col in range(min_col, max_col + 1):
                        col_letter = get_column_letter(col)
                        child_addr = f"{col_letter}{row}"

                        if child_addr == parent_addr:
                            continue

                        if child_addr in cells_dict:
                            cells_dict[child_addr].is_array_formula = True
                            cells_dict[child_addr].parent_array_cell = parent_addr
                        else:
                            cells_dict[child_addr] = CellData(
                                address=child_addr,
                                is_array_formula=True,
                                parent_array_cell=parent_addr
                            )
        # --- END Array Formula Spill Propagation ---
        
        worksheets_data[sheet_name] = WorksheetData(
            name=sheet_name,
            visibility=visibility,
            cells=cells_dict,
            tables=sheet_tables
        )

    # --- OPTIMIZED: Second pass using bounding box checks ---
    for sheet_name, ws_data in worksheets_data.items():
        # Pre-fetch the list of bounding boxes for this sheet to avoid dictionary lookups in the loop
        sheet_bounds = range_references.get(sheet_name, [])
        
        for cell_address, cell in ws_data.cells.items():
            if cell.formula:
                # 1. Check exact O(1) match first
                global_address = f"{sheet_name}!{cell_address}"
                is_used = global_address in exact_referenced_cells
                
                # 2. If not exactly matched, check if it falls inside any bounding box
                if not is_used and sheet_bounds:
                    row_idx, col_idx = coordinate_to_tuple(cell_address)
                    for min_col, min_row, max_col, max_row in sheet_bounds:
                        if (min_row <= row_idx <= max_row) and (min_col <= col_idx <= max_col):
                            is_used = True
                            break # Found a match, no need to check other boxes
                            
                # If it's never referenced exactly or inside a range, it is terminal
                if not is_used:
                    cell.is_terminal = True
    # --------------------------------------------------------

    return worksheets_data

def _has_actual_code(content: str) -> bool:
    """Checks if a VBA module contains anything other than boilerplate attributes."""
    if not content:
        return False
    # Strip out all 'Attribute ...' lines
    code_only = re.sub(r'(?m)^Attribute\s+.*$', '', content)
    # Strip out 'Option Explicit' (often auto-inserted even in empty modules)
    code_only = re.sub(r'(?mi)^\s*Option\s+Explicit\s*$', '', code_only)
    
    # If anything remains besides whitespace, it contains actual code
    return bool(code_only.strip())

def extract_vba_modules(path: Path, wb_formulas) -> list[VBAModule]:
    """Extract VBA macros from the workbook and return them as VBAModule schemas."""
    if VBA_Parser is None:
        print("Warning: oletools not installed. Skipping VBA extraction.")
        return []

    # --- Build the CodeName -> Sheet Name map ---
    codename_to_tab = {}
    for sheet in wb_formulas.worksheets:
        try:
            codename = sheet.sheet_properties.codeName
            if codename:
                codename_to_tab[codename.lower()] = sheet.title
        except AttributeError:
            continue

    modules = []
    try:
        parser = VBA_Parser(str(path))
        if parser.detect_vba_macros():
            for _, _, filename, content in parser.extract_all_macros():
                if isinstance(content, bytes):
                    content = content.decode('utf-8', errors='ignore')
                else:
                    content = str(content) if content is not None else ""
                
                # --- Skip empty/boilerplate modules ---
                if not _has_actual_code(content):
                    continue
                
                safe_filename = str(filename) or "Unknown"
                basename = safe_filename.split('.')[0].lower()  # e.g., "Sheet1.cls" -> "sheet1"
                linked_sheet = codename_to_tab.get(basename)
                
                modules.append(VBAModule(
                    filename=safe_filename,
                    content=content,
                    linked_worksheet=linked_sheet
                ))
    except Exception as e:
        print(f"Warning: Failed to extract VBA modules - {e}")
    finally:
        if 'parser' in locals():
            parser.close()
            
    return modules


def extract_named_ranges(wb) -> list[NamedRange]:
    """
    Extract workbook and worksheet scoped named ranges.
    """
    named_ranges = []
    
    # openpyxl >= 3.1 stores defined names in a dict-like object
    try:
        if hasattr(wb, "defined_names") and hasattr(wb.defined_names, "values"):
            dn_iter = wb.defined_names.values()
        else:
            dn_iter = wb.defined_names.definedName  # openpyxl < 3.1
    except AttributeError:
        return named_ranges

    for dn in dn_iter:
        if isinstance(dn, str): 
            continue  # Safeguard if iteration yields dict keys directly
        
        try:
            name = getattr(dn, "name", None)
            # The formula/reference is usually in 'attr_text' or 'value'
            refers_to = getattr(dn, "attr_text", None) or getattr(dn, "value", None)
            
            if not name or not refers_to:
                continue

            # Determine scope. localSheetId is not None if scoped to a specific sheet.
            localSheetId = getattr(dn, "localSheetId", None)
            if localSheetId is not None:
                try:
                    scope = wb.worksheets[int(localSheetId)].title
                except (IndexError, ValueError, TypeError):
                    scope = f"SheetId_{localSheetId}"
            else:
                scope = "Workbook"

            named_ranges.append(NamedRange(
                name=str(name),
                refers_to=str(refers_to),
                scope=scope
            ))
        except Exception as e:
            print(f"Warning: Failed to parse a named range - {e}")
            continue

    return named_ranges

def extract_tables(sheet) -> list[ExcelTable]:
    """
    Extract structured Excel Data Tables (ListObjects) from a single worksheet.
    """
    tables = []
    
    # openpyxl 3.x stores tables in sheet.tables, accessible via .values()
    try:
        sheet_tables = sheet.tables.values()
    except AttributeError:
        # Fallback for older openpyxl versions
        raw = getattr(sheet, "_tables", None)
        sheet_tables = raw.values() if isinstance(raw, dict) else (raw or [])

    for table in sheet_tables:
        if isinstance(table, str):
            continue  # Safeguard against key leakage
            
        name = getattr(table, "name", None) or getattr(table, "displayName", "UnknownTable")
        ref = getattr(table, "ref", None)
        
        # Extract column headers if available
        columns = []
        if hasattr(table, "tableColumns"):
            for col in table.tableColumns:
                if col.name:
                    columns.append(col.name)

        if name and ref:
            tables.append(ExcelTable(
                name=str(name),
                range_address=str(ref),
                columns=columns
            ))

    return tables

def load_workbooks(path: Path):
    """
    Dual-loads the workbook to extract both structural/formula data 
    and evaluated cached values.
    """
    # Load for structure, formulas, and VBA detection
    wb_formulas = load_workbook(str(path), data_only=False, read_only=False)
    
    # Load strictly for cached cell values
    wb_values = load_workbook(str(path), data_only=True, read_only=False)
    
    return wb_formulas, wb_values

def extract_metadata(path: Path, wb_formulas, vba_modules: list, named_ranges: list) -> WorkbookMetadata:
    """
    Generates the WorkbookMetadata model to provide an immediate summary.
    """
    # Check for hidden sheets
    has_hidden = any(sheet.sheet_state in ('hidden', 'veryHidden') for sheet in wb_formulas.worksheets)
    
    # Check for data tables (ListObjects) across all sheets
    has_tables = any(bool(sheet.tables) for sheet in wb_formulas.worksheets)
    
    # Array formulas require iterating through cells, but we can set this flag 
    # later during the WorksheetData/CellData extraction phase. 
    # Defaulting to False here.
    
    return WorkbookMetadata(
        file_name=path.name,
        file_path=str(path.absolute()),
        file_type=path.suffix.lower(),
        has_macros=len(vba_modules) > 0,
        has_array_formulas=False, # To be updated during cell extraction
        has_data_tables=has_tables,
        has_named_ranges=len(named_ranges) > 0,
        has_hidden_sheets=has_hidden,
        has_external_links=bool(wb_formulas._external_links),
        worksheet_count=len(wb_formulas.worksheets)
    )